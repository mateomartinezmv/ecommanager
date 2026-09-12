#!/usr/bin/env python3
"""Llena la plantilla oficial de carga masiva de Mercado Libre con los 22 SKUs
nuevos de la PI10196.

Entrada:  la plantilla descargada del panel de ML (hojas Manillares,
          Porta Celulares, Espejos; headers en fila 3, datos desde la fila 8).
Precios:  los cargados por el usuario en la columna "Precio oficial" de la
          planilla PI10196-carga-masiva-MELIfin.xlsx.

Sólo se escriben las celdas de datos. Las fórmulas que trae la plantilla
(Cantidad de caracteres, Cargo por venta, Resumen de errores, BUYBOX_FORMULA)
y las validaciones de cada columna quedan intactas.
"""

import shutil

from openpyxl import load_workbook

PLANTILLA = ("/root/.claude/uploads/cc408d9c-ed38-50d7-b709-d354a39d2306/"
             "2faf0bc1-Publicar-09-12-13_21_43.xlsx")
OUT = "/home/user/ecommanager/scripts/meli/PI10196-plantilla-ML-completa.xlsx"

FILA_1 = 8  # primera fila de datos en las tres hojas

# ── Decisiones del usuario ────────────────────────────────────────────────────
MARCA        = "Genérico"
MODELO       = "Universal"
GARANTIA     = ("Garantía del vendedor", 90, "días")
FORMA_ENVIO  = "Mercado Envíos | Mercado Envíos Flex"
RETIRO       = "Acepto"
CONDICION    = "Nuevo"
MONEDA       = "$"
UMBRAL_GRATIS = 900   # envío gratis sólo arriba de $900


def costo_envio(precio):
    return "Ofreces envío gratis" if precio > UMBRAL_GRATIS else "A cargo del comprador"


# ── Descripciones ─────────────────────────────────────────────────────────────
ROSCA = {"F": "horario", "R": "antihorario"}


def ext_roscas(sku):
    _, _, a, b, _ = sku.split("-")
    return f"{a[:-1]}mm {ROSCA[a[-1]]}", f"{b[:-1]}mm {ROSCA[b[-1]]}"


DESC_SOPORTE = (
    "Soporte universal para celular apto para moto, con anclaje {anclaje}.\n\n"
    "- Sujeción ajustable: se adapta a celulares de 4,7 a 7 pulgadas.\n"
    "- Rotación 360 grados: podés usarlo en vertical u horizontal.\n"
    "- Construcción en ABS reforzado con gomas antivibración.\n"
    "- Instalación sin herramientas especiales.\n\n"
    "Producto nuevo. Garantía del vendedor: 90 días."
)

DESC_ESPEJO_CIR = (
    "Par de espejos de puño circulares para moto, universales.\n\n"
    "- Montaje en manillar de 7/8 pulgadas (22mm).\n"
    "- Cuerpo de aluminio con terminación negra.\n"
    "- Brazo y cabezal regulables para ajustar el ángulo de visión.\n"
    "- Se venden por par (izquierdo y derecho).\n\n"
    "Producto nuevo. Garantía del vendedor: 90 días."
)

DESC_EXT = (
    "Extensor y adaptador de rosca para espejo de moto.\n\n"
    "- Rosca interior: {interior}.\n"
    "- Rosca exterior: {exterior}.\n"
    "- Permite montar espejos con rosca distinta a la del soporte original y "
    "ganar altura para mejorar el campo de visión.\n"
    "- Acero con terminación negra.\n"
    "- Se vende por unidad.\n\n"
    "Producto nuevo. Garantía del vendedor: 90 días."
)

DESC_MANILLAR = (
    "{tipo}, universal para moto.\n\n"
    "- Medida universal 7/8 pulgadas (22mm) de diámetro.\n"
    "- Aluminio de alta resistencia.\n"
    "- Compatible con puños, espejos y comandos estándar.\n"
    "- Se vende por unidad.{extra}\n\n"
    "Producto nuevo. Garantía del vendedor: 90 días."
)

# ═══════════════════════════════════════════════════════════════════════════════
# Datos por hoja
# ═══════════════════════════════════════════════════════════════════════════════

# ── Manillares ────────────────────────────────────────────────────────────────
# Los dos V2 comparten título y precio: ML los toma como una sola publicación
# con variante de color. El V3 es una publicación aparte.
TITULO_V2 = "Manillar Moto Universal 7/8 22mm Alto Antideslizante"

MANILLARES = [
    {"titulo": TITULO_V2, "color": "Negro", "sku": "MAN-78-V2-NEG-001",
     "stock": 5, "precio": 690,
     "desc": DESC_MANILLAR.format(
         tipo="Manillar alto con superficie antideslizante",
         extra=" Disponible en negro y plateado.")},
    {"titulo": TITULO_V2, "color": "Plateado", "sku": "MAN-78-V2-PLA-001",
     "stock": 5, "precio": 690,
     "desc": DESC_MANILLAR.format(
         tipo="Manillar alto con superficie antideslizante",
         extra=" Disponible en negro y plateado.")},
    {"titulo": "Manillar Café Racer Universal 7/8 22mm Recto Plateado Cromado",
     "color": "Plateado", "sku": "MAN-78-V3-PLA-001",
     "stock": 5, "precio": 890,
     "desc": DESC_MANILLAR.format(
         tipo="Manillar tipo Café Racer, recto, con terminación cromada",
         extra="")},
]

# ── Porta Celulares ───────────────────────────────────────────────────────────
PORTA_CELULARES = [
    {"titulo": "Soporte Celular Moto Universal Manillar Negro Ajustable",
     "color": "Negro", "sku": "SOP-CEL-MAN-001", "stock": 100, "precio": 490,
     "desc": DESC_SOPORTE.format(anclaje="al manillar")},
    {"titulo": "Soporte Celular Moto Universal Espejo Negro Ajustable",
     "color": "Negro", "sku": "SOP-CEL-ESP-001", "stock": 100, "precio": 490,
     "desc": DESC_SOPORTE.format(anclaje="a la base del espejo")},
]

# ── Espejos (incluye los extensores, por pedido del usuario) ──────────────────
EXTENSORES = [
    ("EXT-ESP-8F-8F-001",   "Extensor Espejo Moto Rosca 8mm Int/Ext Horario Universal"),
    ("EXT-ESP-8F-8R-001",   "Extensor Espejo Moto Rosca 8mm Int Hor / Ext Antihor"),
    ("EXT-ESP-8R-8F-001",   "Extensor Espejo Moto Rosca 8mm Int Antihor / Ext Hor"),
    ("EXT-ESP-8R-8R-001",   "Extensor Espejo Moto Rosca 8mm Int/Ext Antihorario Univ"),
    ("EXT-ESP-8F-10F-001",  "Extensor Espejo Moto Rosca 8mm Int Hor / 10mm Ext Hor"),
    ("EXT-ESP-8R-10F-001",  "Extensor Espejo Moto Rosca 8mm Int Antihor / 10mm Ext"),
    ("EXT-ESP-8F-10R-001",  "Extensor Espejo Moto Rosca 8mm Int Hor / 10mm Ext Antihor"),
    ("EXT-ESP-8R-10R-001",  "Extensor Espejo Moto Rosca 8mm Int/10mm Ext Antihorario"),
    ("EXT-ESP-10F-10F-001", "Extensor Espejo Moto Rosca 10mm Int/Ext Horario Universal"),
    ("EXT-ESP-10R-10F-001", "Extensor Espejo Moto Rosca 10mm Int Antihor / Ext Hor"),
    ("EXT-ESP-10F-10R-001", "Extensor Espejo Moto Rosca 10mm Int Hor / Ext Antihor"),
    ("EXT-ESP-10R-10R-001", "Extensor Espejo Moto Rosca 10mm Int/Ext Antihorario Univ"),
    ("EXT-ESP-10F-8F-001",  "Extensor Espejo Moto Rosca 10mm Int Hor / 8mm Ext Hor"),
    ("EXT-ESP-10F-8R-001",  "Extensor Espejo Moto Rosca 10mm Int Hor / 8mm Ext Antihor"),
    ("EXT-ESP-10R-8F-001",  "Extensor Espejo Moto Rosca 10mm Int Antihor / 8mm Ext"),
    ("EXT-ESP-10R-8R-001",  "Extensor Espejo Moto Rosca 10mm Int/8mm Ext Antihorario"),
]

ESPEJOS = [
    {"titulo": "Espejos de Puño Circulares Moto Universal Manillar 7/8 Negro",
     "sku": "ESP-CIR-001", "stock": 10, "precio": 1090, "desc": DESC_ESPEJO_CIR,
     "lado": "Ambos lados", "ajustable": "Sí"},
]
for sku, titulo in EXTENSORES:
    interior, exterior = ext_roscas(sku)
    ESPEJOS.append({
        "titulo": titulo, "sku": sku, "stock": 10, "precio": 100,
        "desc": DESC_EXT.format(interior=interior, exterior=exterior),
        # El sentido de la rosca es justamente lo que define el lado, y no
        # tengo el mapeo rosca->lado de cada moto: se deja sin completar.
        "lado": None, "ajustable": "No",
    })

assert len(MANILLARES) + len(PORTA_CELULARES) + len(ESPEJOS) == 22

# ═══════════════════════════════════════════════════════════════════════════════
# Escritura
# ═══════════════════════════════════════════════════════════════════════════════
shutil.copy(PLANTILLA, OUT)
wb = load_workbook(OUT)


def poner(ws, fila, col, valor):
    if valor is not None:
        ws.cell(row=fila, column=col, value=valor)


# ── Manillares ────────────────────────────────────────────────────────────────
ws = wb["Manillares"]
for i, it in enumerate(MANILLARES):
    r = FILA_1 + i
    poner(ws, r, 1,  it["titulo"])
    poner(ws, r, 3,  CONDICION)
    poner(ws, r, 4,  it["color"])
    # 5 Fotos: lo completa el usuario desde el Gestor de fotos de ML
    poner(ws, r, 6,  it["sku"])
    poner(ws, r, 7,  it["stock"])
    poner(ws, r, 8,  it["precio"])
    poner(ws, r, 9,  MONEDA)
    poner(ws, r, 10, it["desc"])
    poner(ws, r, 12, FORMA_ENVIO)
    poner(ws, r, 13, costo_envio(it["precio"]))
    poner(ws, r, 14, RETIRO)
    poner(ws, r, 15, GARANTIA[0])
    poner(ws, r, 16, GARANTIA[1])
    poner(ws, r, 17, GARANTIA[2])
    poner(ws, r, 18, MARCA)
    poner(ws, r, 19, MODELO)
    poner(ws, r, 20, it["sku"])          # Número de pieza
    poner(ws, r, 21, "Aluminio")         # Material
    # 22-27 Largo/Ancho/Altura: sin medidas del proveedor, se dejan vacíos
    # 28 Estilo de conducción en moto: opcional, se deja sin completar

# ── Porta Celulares ───────────────────────────────────────────────────────────
ws = wb["Porta Celulares"]
for i, it in enumerate(PORTA_CELULARES):
    r = FILA_1 + i
    # 1 Código de catálogo ML: vacío, no es publicación de catálogo
    poner(ws, r, 2,  it["titulo"])
    poner(ws, r, 4,  CONDICION)
    # 5 Código universal de producto: opcional, sin EAN
    poner(ws, r, 6,  it["color"])
    # 7 Fotos: trae una fórmula de catálogo; la completa el usuario
    poner(ws, r, 8,  it["sku"])
    poner(ws, r, 9,  it["stock"])
    poner(ws, r, 10, it["precio"])
    poner(ws, r, 11, MONEDA)
    poner(ws, r, 12, it["desc"])
    poner(ws, r, 14, FORMA_ENVIO)
    poner(ws, r, 15, costo_envio(it["precio"]))
    poner(ws, r, 16, RETIRO)
    poner(ws, r, 17, GARANTIA[0])
    poner(ws, r, 18, GARANTIA[1])
    poner(ws, r, 19, GARANTIA[2])
    poner(ws, r, 20, MARCA)
    poner(ws, r, 21, MODELO)
    # 22 Tipo de soporte: sólo acepta "De pared" o "Repisa"; ninguno aplica
    poner(ws, r, 23, "Plástico")         # Material
    poner(ws, r, 24, "Sí")               # Ajustable

# ── Espejos ───────────────────────────────────────────────────────────────────
ws = wb["Espejos"]
for i, it in enumerate(ESPEJOS):
    r = FILA_1 + i
    poner(ws, r, 1,  it["titulo"])
    poner(ws, r, 3,  CONDICION)
    poner(ws, r, 4,  "El producto no tiene código registrado")
    # 5 Fotos: lo completa el usuario
    poner(ws, r, 6,  it["sku"])
    poner(ws, r, 7,  it["stock"])
    poner(ws, r, 8,  it["precio"])
    poner(ws, r, 9,  MONEDA)
    poner(ws, r, 10, it["desc"])
    poner(ws, r, 12, FORMA_ENVIO)
    poner(ws, r, 13, costo_envio(it["precio"]))
    poner(ws, r, 14, RETIRO)
    poner(ws, r, 15, GARANTIA[0])
    poner(ws, r, 16, GARANTIA[1])
    poner(ws, r, 17, GARANTIA[2])
    poner(ws, r, 18, MARCA)
    poner(ws, r, 19, it["sku"])          # Número de pieza
    poner(ws, r, 20, it["lado"])
    # 21-22 Largo: sin medidas del proveedor
    poner(ws, r, 23, "Mate")             # Acabado de la superficie
    poner(ws, r, 24, it["ajustable"])
    poner(ws, r, 25, "China")            # Origen
    # 26 Estilo de conducción en moto: opcional, se deja sin completar

wb.save(OUT)
print("OK ->", OUT)
print("Manillares:", len(MANILLARES), "| Porta Celulares:", len(PORTA_CELULARES),
      "| Espejos:", len(ESPEJOS))
